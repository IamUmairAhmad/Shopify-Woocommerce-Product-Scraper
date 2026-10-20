"""Tests for: Excel export, several collections, specific products, keyword search, product-link column,
reviews and the product details view. Run with `python3 -I tests/test_features.py`."""

import csv
import io
import os
import subprocess
import sys
import time
import zipfile

os.environ["ALLOW_PRIVATE_HOSTS"] = "1"
os.environ["SCRAPE_DELAY_SCALE"] = "0"
os.environ["NO_PROXY"] = "127.0.0.1,localhost"
os.environ["no_proxy"] = "127.0.0.1,localhost"
os.environ["RATE_LIMIT_PER_MINUTE"] = "10000"

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import openpyxl  # noqa: E402
import requests  # noqa: E402

import app as exporter  # noqa: E402
import scraper  # noqa: E402

failures = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail and not cond else ""))
    if not cond:
        failures.append(name)


def start(script, port, **env):
    e = dict(os.environ, **{k: str(v) for k, v in env.items()})
    p = subprocess.Popen([sys.executable, os.path.join(HERE, script), str(port)], env=e,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(100):
        try:
            requests.get(f"http://127.0.0.1:{port}/", timeout=1)
            break
        except requests.exceptions.RequestException:
            time.sleep(0.1)
    return p


def run_job(c, payload, timeout=120):
    r = c.post("/api/search", json=payload)
    if r.status_code != 200:
        return r, None
    jid = r.json["job_id"]
    end = time.time() + timeout
    while time.time() < end:
        j = c.get(f"/api/job/{jid}").json
        if j["status"] != "running":
            return r, j
        time.sleep(0.05)
    raise TimeoutError


def rows_of(resp):
    return list(csv.DictReader(io.StringIO(resp.get_data(as_text=True))))


def export(c, token, fmt, **extra):
    return c.post("/api/export", json={"token": token, "format": fmt, **extra})


def count(j):
    return j["result"]["count"] if j and j.get("result") else None


def unit_checks():
    ph = scraper.parse_handles
    check("handles: urls, slugs, json links", ph("https://a.com/products/x-1?variant=2\nhttps://a.com/product/y-2/, z-3 a.com/products/w-4.json")
          == ["x-1", "y-2", "z-3", "w-4"], str(ph("https://a.com/products/x-1?variant=2\nhttps://a.com/product/y-2/, z-3 a.com/products/w-4.json")))
    check("handles: duplicates and junk dropped", ph("a a $bad https://a.com/products/a") == ["a"], str(ph("a a $bad https://a.com/products/a")))
    check("handles: empty", ph("  \n ") == [])
    check("search: all words must match", scraper.product_matches({"title": "Red Running Shoe", "tags": ["sale"]}, "running red")
          and not scraper.product_matches({"title": "Red Shoe"}, "red boot"))


def main():
    unit_checks()
    procs = [
        start("mock_store.py", 5081, MOCK_PRODUCTS=120),
        start("mock_woo.py", 5082),
        start("mock_woo.py", 5083, MOCK_MODE="plain", MOCK_BULK_VARIATIONS=0, MOCK_MAX_PER_PAGE=50),
        start("mock_woo.py", 5084, MOCK_REVIEWS=0),
    ]
    SH, WO, PL, NR = (f"http://127.0.0.1:{p}" for p in (5081, 5082, 5083, 5084))
    try:
        c = exporter.app.test_client()

        # ---------------- WooCommerce: several categories, search, specific products
        _, j = run_job(c, {"store": WO, "scope": "collections", "collections": ["hoodies", "clothing"]})
        check("woo: overlapping categories are de-duplicated", count(j) == 10, str(j and (j.get("error") or count(j))))
        _, j = run_job(c, {"store": WO, "scope": "collections", "collections": ["hoodies", "accessories"]})
        check("woo: two categories combined", count(j) == 245, str(count(j)))
        _, j = run_job(c, {"store": WO, "scope": "search", "query": "hoodie"})
        check("woo: keyword search (on the store)", count(j) == 10, str(j and (j.get("error") or count(j))))
        _, j = run_job(c, {"store": WO, "scope": "search", "query": "beanie 1"})
        check("woo: search matches decoded names", count(j) == 111, str(count(j)))
        _, j = run_job(c, {"store": WO, "scope": "search", "query": "zzzz"})
        check("woo: search with no match reports it", j and j["status"] == "error" and "No products" in j["error"], str(j and j.get("error")))
        text = f"{WO}/product/hoodie-2/\nbeanie-7, nope-xyz\n{WO}/product/hoodie-2/"
        _, j = run_job(c, {"store": WO, "scope": "products", "products_text": text})
        check("woo: specific products by link/slug", count(j) == 2, str(j and (j.get("error") or count(j))))
        check("woo: missing product is reported", j and any("nope-xyz" in w for w in j["warnings"]), str(j and j["warnings"]))
        _, j = run_job(c, {"store": PL, "scope": "products", "products_text": "hoodie-3"})
        check("woo (plain permalinks): specific product", count(j) == 1, str(j and (j.get("error") or count(j))))
        r, _ = run_job(c, {"store": WO, "scope": "products", "products_text": "  "})
        check("specific products: empty input is rejected", r.status_code == 400)
        r, _ = run_job(c, {"store": WO, "scope": "search", "query": ""})
        check("search: empty keyword is rejected", r.status_code == 400)

        # ---------------- WooCommerce: reviews
        _, j = run_job(c, {"store": WO, "scope": "all", "reviews": True})
        check("woo: reviews job done", j and j["status"] == "done", str(j and j.get("error")))
        tok = j["result"]["token"]
        rv = rows_of(export(c, tok, "reviews"))
        check("woo: 76 reviews exported", len(rv) == 76, str(len(rv)))
        first = rv[0]
        check("woo: review text is clean", first["Review"] == "Great & warm, review 1." and first["Reviewer"] == "Reviewer 1", str(first))
        check("woo: review links to its product", first["Product Handle"] and first["Product URL"].startswith("http") and first["Product Title"])
        check("woo: verified flag", {r["Verified"] for r in rv} == {"TRUE", "FALSE"})
        js = export(c, tok, "json").get_json(force=True)
        check("woo: json carries reviews", sum(len(i.get("reviews", [])) for i in js) == 76)
        _, j2 = run_job(c, {"store": PL, "scope": "all", "reviews": True})
        check("woo (page cap 50): all reviews collected", j2 and j2["status"] == "done" and sum(1 for _ in rows_of(export(c, j2["result"]["token"], "reviews"))) == 76)
        _, j3 = run_job(c, {"store": NR, "scope": "all", "reviews": True})
        check("woo: store without reviews api warns", j3 and j3["status"] == "done" and any("reviews" in w for w in j3["warnings"]), str(j3 and j3["warnings"]))
        r = export(c, j3["result"]["token"], "reviews")
        check("reviews export without reviews is a clear error", r.status_code == 400 and "reviews" in r.get_json()["error"].lower())

        # ---------------- Excel, link column, details view
        full = rows_of(export(c, tok, "full"))
        x = export(c, tok, "xlsx")
        check("xlsx: served with the right type", x.status_code == 200 and "spreadsheetml" in x.mimetype)
        wb = openpyxl.load_workbook(io.BytesIO(x.data))
        check("xlsx: Products and Reviews sheets", wb.sheetnames == ["Products", "Reviews"], str(wb.sheetnames))
        ws = wb["Products"]
        hdr = [cell.value for cell in ws[1]]
        check("xlsx: same rows as the full-details csv", ws.max_row - 1 == len(full), f"{ws.max_row - 1} vs {len(full)}")
        check("xlsx: same columns", hdr == list(full[0].keys()))
        price_idx = hdr.index("Price")
        check("xlsx: prices are real numbers", isinstance(ws.cell(2, price_idx + 1).value, (int, float)), repr(ws.cell(2, price_idx + 1).value))
        check("xlsx: reviews sheet rows", wb["Reviews"].max_row - 1 == 76)
        check("xlsx: opens cleanly as a zip", zipfile.ZipFile(io.BytesIO(x.data)).testzip() is None)

        plain = rows_of(export(c, tok, "woocommerce"))
        check("woo csv: no link column by default", "Product URL" not in plain[0])
        linked = rows_of(export(c, tok, "woocommerce", include_link=True))
        check("woo csv: link column on request", linked[0].get("Product URL", "").startswith("http"), str(linked[0].get("Product URL")))
        check("woo csv: variations carry no link", all(not r.get("Product URL") for r in linked if r["Type"] == "variation"))
        sl = rows_of(export(c, tok, "shopify", include_link=True))
        check("shopify csv: link column on request", "Product URL" in sl[0] and sl[0]["Product URL"].startswith("http"))

        r = c.post("/api/product", json={"token": tok, "id": 1001})
        d = r.get_json()
        check("details view: product 1001", r.status_code == 200 and d["title"] == "Hoodie 1" and len(d["variants"]) == 4, str(d)[:200])
        check("details view: reviews and attributes", d["review_total"] == 3 and any(a["name"] == "Material" for a in d["attributes"]) and d["url"].startswith("http"))
        check("details view: unknown product", c.post("/api/product", json={"token": tok, "id": 424242}).status_code == 404)

        # ---------------- Shopify
        _, j = run_job(c, {"store": SH, "scope": "collections", "collections": ["shoes", "sale"]})
        check("shopify: overlapping collections are de-duplicated", count(j) == 40, str(j and (j.get("error") or count(j))))
        _, j = run_job(c, {"store": SH, "scope": "collections", "collections": ["shoes", "socks"], "include_collections": False})
        check("shopify: two collections combined", count(j) == 120, str(count(j)))
        _, j = run_job(c, {"store": SH, "scope": "search", "query": "pro edition", "include_collections": False})
        check("shopify: keyword search", count(j) == 40, str(j and (j.get("error") or count(j))))
        text = f"{SH}/products/runner-3\nsock-1 {SH}/collections/shoes/products/sock-2 missing-one"
        _, j = run_job(c, {"store": SH, "scope": "products", "products_text": text, "include_collections": False})
        check("shopify: specific products by link/handle", count(j) == 3, str(j and (j.get("error") or count(j))))
        check("shopify: missing product is reported", j and any("missing-one" in w for w in j["warnings"]), str(j and j["warnings"]))
        _, j = run_job(c, {"store": SH, "scope": "all", "reviews": True, "include_collections": False})
        check("shopify: reviews job done", j and j["status"] == "done", str(j and j.get("error")))
        sk = j["result"]["token"]
        rv = rows_of(export(c, sk, "reviews"))
        with_reviews = [i for i in range(1, 121) if (1000 + i) % 7 == 0 and (1000 + i) % 50 != 0]
        check("shopify: reviews read from structured data", len(rv) == 2 * len(with_reviews), f"{len(rv)} vs {2 * len(with_reviews)}")
        check("shopify: review fields", rv and rv[0]["Reviewer"] == "Sam" and rv[0]["Rating"] == "5" and rv[0]["Review"] == "Fits well & looks great."
              and rv[0]["Title"] == "Love them", str(rv[:1]))
        _, j = run_job(c, {"store": SH, "scope": "all", "include_collections": False})
        check("shopify: no reviews requested -> none exported", export(c, j["result"]["token"], "reviews").status_code == 400)
        sx = openpyxl.load_workbook(io.BytesIO(export(c, sk, "xlsx").data))
        check("shopify xlsx has a Reviews sheet", "Reviews" in sx.sheetnames)
        sl = rows_of(export(c, sk, "shopify", include_link=True))
        n_products = j["result"]["count"]
        check("shopify csv: link only on each product's first row",
              sum(1 for r in sl if r["Product URL"]) == n_products and all("/products/" in r["Product URL"] for r in sl if r["Product URL"]),
              f"{sum(1 for r in sl if r['Product URL'])} vs {n_products}")
    finally:
        for p in procs:
            p.terminate()
    print()
    print("ALL PASSED" if not failures else f"{len(failures)} FAILED: {failures}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
