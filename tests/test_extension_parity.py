"""Checks that the Chrome extension's JavaScript scraper and exporters produce the same files as the
Python tool, against the mock stores. Run with `python3 -I tests/test_extension_parity.py` (needs Node 20+)."""

import json
import os
import shutil
import subprocess
import sys
import tempfile
import time

os.environ["ALLOW_PRIVATE_HOSTS"] = "1"
os.environ["SCRAPE_DELAY_SCALE"] = "0"
os.environ["NO_PROXY"] = "127.0.0.1,localhost"
os.environ["no_proxy"] = "127.0.0.1,localhost"
os.environ.pop("HTTP_PROXY", None); os.environ.pop("HTTPS_PROXY", None)
os.environ.pop("http_proxy", None); os.environ.pop("https_proxy", None)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
import requests  # noqa: E402

failures = []


def check(name, cond, detail=""):
    print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail and not cond else ""))
    if not cond:
        failures.append(name)


def start(script, port, **env):
    e = dict(os.environ, **{k: str(v) for k, v in env.items()})
    p = subprocess.Popen([sys.executable, os.path.join(HERE, script), str(port)], env=e,
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    for _ in range(100):
        try:
            requests.get(f"http://127.0.0.1:{port}/", timeout=1)
            break
        except requests.exceptions.RequestException:
            time.sleep(0.1)
    return p


def both(store, extra=()):
    tmp = tempfile.mkdtemp()
    py, js = os.path.join(tmp, "py"), os.path.join(tmp, "js")
    r1 = subprocess.run([sys.executable, os.path.join(ROOT, "cli.py"), store, "--out", py, *extra],
                        capture_output=True, text=True, env=os.environ)
    r2 = subprocess.run(["node", os.path.join(ROOT, "chrome-extension", "test", "run-cli.mjs"), store, js, *extra],
                        capture_output=True, text=True, env=os.environ)
    return py, js, r1, r2


def same_xlsx(a, b):
    """Zip bytes differ between writers, so compare what Excel would show: sheet names and cell values."""
    import openpyxl
    wa, wb = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    if wa.sheetnames != wb.sheetnames:
        return False, f"sheets {wa.sheetnames} vs {wb.sheetnames}"
    for name in wa.sheetnames:
        ra = [list(r) for r in wa[name].iter_rows(values_only=True)]
        rb = [list(r) for r in wb[name].iter_rows(values_only=True)]
        if ra != rb:
            for i, (x, y) in enumerate(zip(ra, rb)):
                if x != y:
                    return False, f"{name} row {i + 1}\n  py: {x[:12]}\n  js: {y[:12]}"
            return False, f"{name}: {len(ra)} vs {len(rb)} rows"
    return True, ""


def compare(label, store, extra=()):
    py, js, r1, r2 = both(store, extra)
    if r1.returncode or r2.returncode:
        check(label + ": both runs succeed", False, (r1.stderr + r1.stdout)[-300:] + (r2.stderr + r2.stdout)[-300:])
        return
    names = sorted(os.listdir(py))
    check(label + ": same files written", names == sorted(os.listdir(js)), f"{names} vs {sorted(os.listdir(js))}")
    for n in names:
        if n.endswith(".xlsx"):
            a = b = ""
        else:
            a = open(os.path.join(py, n), encoding="utf-8", newline="").read()
            b = open(os.path.join(js, n), encoding="utf-8", newline="").read()
        if n.endswith(".xlsx"):
            ok, detail = same_xlsx(os.path.join(py, n), os.path.join(js, n))
            check(f"{label}: {n.split('-', 1)[1]} same cells", ok, detail)
            continue
        if n.endswith(".json"):
            ok = json.loads(a) == json.loads(b)
        else:
            ok = a == b
        detail = ""
        if not ok and not n.endswith(".json"):
            la, lb = a.split("\n"), b.split("\n")
            for i, (x, y) in enumerate(zip(la, lb)):
                if x != y:
                    detail = f"line {i + 1}\n  py: {x[:300]}\n  js: {y[:300]}"
                    break
            else:
                detail = f"length {len(la)} vs {len(lb)}"
        check(f"{label}: {n.split('-', 1)[1]} identical", ok, detail)
    shutil.rmtree(os.path.dirname(py), ignore_errors=True)


def main():
    procs = [
        start("mock_store.py", 5071, MOCK_PRODUCTS=600),
        start("mock_woo.py", 5072),
        start("mock_woo.py", 5073, MOCK_MODE="plain", MOCK_BULK_VARIATIONS=0, MOCK_MAX_PER_PAGE=50),
    ]
    try:
        compare("shopify", "http://127.0.0.1:5071")
        compare("shopify deep", "http://127.0.0.1:5071", ["--deep"])
        compare("woo", "http://127.0.0.1:5072")
        compare("woo deep", "http://127.0.0.1:5072", ["--deep"])
        compare("woo plain permalinks", "http://127.0.0.1:5073")
        compare("woo category", "http://127.0.0.1:5072/product-category/clothing/hoodies/")
        compare("woo no variations", "http://127.0.0.1:5072", ["--no-variations"])
        compare("shopify collections", "http://127.0.0.1:5071", ["--collections", "shoes,sale", "--link"])
        compare("shopify search", "http://127.0.0.1:5071", ["--search", "runner 1"])
        compare("shopify products", "http://127.0.0.1:5071",
                ["--products", "http://127.0.0.1:5071/products/runner-3", "sock-2", "nope-x"])
        compare("shopify reviews", "http://127.0.0.1:5071", ["--collections", "shoes", "--reviews", "--link"])
        compare("woo reviews", "http://127.0.0.1:5072", ["--reviews", "--link"])
        compare("woo collections", "http://127.0.0.1:5072", ["--collections", "hoodies,accessories"])
        compare("woo search", "http://127.0.0.1:5072", ["--search", "hoodie"])
        compare("woo products", "http://127.0.0.1:5072", ["--products", "hoodie-1", "hoodie-2"])
    finally:
        for p in procs:
            p.terminate()
    print()
    print("ALL PASSED" if not failures else f"{len(failures)} FAILED: {failures}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
